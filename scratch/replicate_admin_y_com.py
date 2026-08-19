import openpyxl
from copy import copy
import re
from openpyxl.chart import LineChart, AreaChart3D, BarChart, Reference

# Load workbooks
src_wb = openpyxl.load_workbook("SGC.REG-004 Seguimiento de OC 2026.xlsx")
dest_wb = openpyxl.load_workbook("reporte_mensual (63).xlsx")

# Check sheet names
if "Admin y Com" in dest_wb.sheetnames:
    del dest_wb["Admin y Com"]

src_ws = src_wb["Admin y Com"]
dest_ws = dest_wb.create_sheet("Admin y Com")
dest_ws.views.sheetView[0].showGridLines = True

# 1. Copy values, formulas, and styles
def translate_formula(formula):
    if not formula or not isinstance(formula, str):
        return formula
    if not formula.startswith("="):
        return formula
        
    # Replace external workbook links
    formula = formula.replace("'[1] 2011-2026'", "'2012 - 2026'")
    formula = formula.replace("' 2011-2026'", "'2012 - 2026'")
    formula = formula.replace("=[1]Objetivos", "=Objetivos")
    formula = formula.replace("'[1]Objetivos'", "'Objetivos'")
    
    # Translate row/column references for '2012 - 2026' sheet
    # M198 -> C147, N198 -> D147, O198 -> E147, P198 -> F147, L198 -> B147
    col_map = {'M': 'C', 'N': 'D', 'O': 'E', 'P': 'F', 'L': 'B'}
    
    def repl_hist(m):
        col = m.group(1)
        row = int(m.group(2))
        new_col = col_map.get(col, col)
        new_row = row - 51
        return f"'2012 - 2026'!{new_col}{new_row}"
        
    formula = re.sub(r"'2012 - 2026'!\$?([L-P])\$?(\d+)", repl_hist, formula)
    
    # Translate Objetivos references: Objetivos!D2 -> Objetivos!D3
    def repl_obj(m):
        col = m.group(1)
        row = int(m.group(2))
        return f"Objetivos!{col}{row + 1}"
        
    formula = re.sub(r"Objetivos!\$?([A-Z])\$?(\d+)", repl_obj, formula)
    
    return formula

# Copy cells
for r in range(1, src_ws.max_row + 1):
    for c in range(1, src_ws.max_column + 1):
        src_cell = src_ws.cell(row=r, column=c)
        dest_cell = dest_ws.cell(row=r, column=c)
        
        # Copy and translate value/formula
        val = src_cell.value
        if isinstance(val, str) and val.startswith("="):
            dest_cell.value = translate_formula(val)
        else:
            dest_cell.value = val
            
        # Copy style
        if src_cell.has_style:
            dest_cell.font = copy(src_cell.font)
            dest_cell.fill = copy(src_cell.fill)
            dest_cell.border = copy(src_cell.border)
            dest_cell.alignment = copy(src_cell.alignment)
            dest_cell.number_format = src_cell.number_format

# Copy merged cells
for merged_range in src_ws.merged_cells.ranges:
    dest_ws.merge_cells(str(merged_range))

# Copy column dimensions
for col_letter, col_dim in src_ws.column_dimensions.items():
    dest_ws.column_dimensions[col_letter].width = col_dim.width

# Copy row dimensions
for row_num, row_dim in src_ws.row_dimensions.items():
    dest_ws.row_dimensions[row_num].height = row_dim.height

# 2. Add the four charts back
# Chart 1: LineChart 'Ordenes de compra vs Facturado'
c1 = LineChart()
c1.title = "Ordenes de compra vs Facturado"
c1.style = 13
c1.width = 15
c1.height = 7.5
data_c1 = Reference(dest_ws, min_col=6, min_row=2, max_col=6, max_row=14) # Monto OC
data_c1_2 = Reference(dest_ws, min_col=11, min_row=2, max_col=11, max_row=14) # Monto Fact
cats_c1 = Reference(dest_ws, min_col=1, min_row=3, max_row=14)

c1.add_data(data_c1, titles_from_data=True)
c1.add_data(data_c1_2, titles_from_data=True)
c1.set_categories(cats_c1)
dest_ws.add_chart(c1, "L1")

# Chart 2: AreaChart3D 'Ordenes de compra VS Facturado'
c2 = AreaChart3D()
c2.title = "Ordenes de compra VS Facturado"
c2.style = 13
c2.width = 15
c2.height = 7.5
c2.add_data(data_c1, titles_from_data=True)
c2.add_data(data_c1_2, titles_from_data=True)
c2.set_categories(cats_c1)
dest_ws.add_chart(c2, "L16")

# Chart 3: BarChart '%OC VS %FAC' (horizontal bar chart)
c3 = BarChart()
c3.type = "bar" # horizontal
c3.grouping = "clustered"
c3.title = "%OC VS %FAC"
c3.width = 15
c3.height = 7.5
data_c3 = Reference(dest_ws, min_col=5, min_row=25, max_row=26)
cats_c3 = Reference(dest_ws, min_col=2, min_row=25, max_row=26)
c3.add_data(data_c3, titles_from_data=False)
c3.set_categories(cats_c3)
c3.legend = None
dest_ws.add_chart(c3, "A28")

# Chart 4: BarChart 'Ordenes de compra VS Facturado' (vertical column chart)
c4 = BarChart()
c4.type = "col" # vertical
c4.grouping = "clustered"
c4.title = "Ordenes de compra VS Facturado"
c4.width = 15
c4.height = 7.5
c4.add_data(data_c1, titles_from_data=True)
c4.add_data(data_c1_2, titles_from_data=True)
c4.set_categories(cats_c1)
dest_ws.add_chart(c4, "K32")

# Save output
dest_wb.save("reporte_mensual (63).xlsx")
print("Replication script completed successfully for reporte_mensual (63).xlsx!")
