import openpyxl

wb = openpyxl.load_workbook("scratch/test_generated_report.xlsx", data_only=False)
if "Admin y Com" in wb.sheetnames:
    ws = wb["Admin y Com"]
    print("--- Admin y Com Sheet Verification (Formulas) ---")
    print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
    print(f"Charts in sheet: {len(ws._charts)}")
    
    # Check rows 1 to 21
    for r in range(1, 22):
        row_vals = []
        for c in range(1, 12):
            cell = ws.cell(row=r, column=c)
            row_vals.append(cell.value)
        print(f"Row {r:02d}: {row_vals}")

    # Check rows 25 to 27
    print("\n--- Rows 25 to 27 ---")
    for r in range(25, 28):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 6)]
        print(f"Row {r:02d}: {row_vals}")

    # Check rows 49 to 54
    print("\n--- Rows 49 to 54 ---")
    for r in range(49, 55):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 7)]
        print(f"Row {r:02d}: {row_vals}")

    # Check evaluated values
    wb_val = openpyxl.load_workbook("scratch/test_generated_report.xlsx", data_only=True)
    ws_val = wb_val["Admin y Com"]
    print("\n--- Admin y Com Sheet Verification (Values) ---")
    for r in range(1, 22):
        row_vals = []
        for c in range(1, 12):
            cell = ws_val.cell(row=r, column=c)
            row_vals.append(cell.value)
        print(f"Row {r:02d}: {row_vals}")

    # Check charts
    print("\n--- Charts Verification ---")
    for idx, chart in enumerate(ws._charts):
        print(f"Chart {idx+1} Class: {chart.__class__.__name__} | Title: {chart.title.text if chart.title else 'No Title'} | Anchor: {chart.anchor}")
        print("  Series:")
        for s_idx, s in enumerate(chart.series):
            print(f"    s{s_idx} Val Ref: {s.val.numRef.f if getattr(s.val, 'numRef', None) else 'No numRef'}")
            print(f"    s{s_idx} Tx Ref: {s.tx.strRef.f if getattr(s.tx, 'strRef', None) else 'No strRef'}")
else:
    print("Error: Admin y Com sheet not found in generated report!")
