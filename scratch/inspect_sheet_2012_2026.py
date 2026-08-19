import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (22).xlsx", data_only=False)
ws = wb["2012 - 2026"]
print(f"Sheet: {ws.title}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
print(f"Number of charts: {len(ws._charts)}")

for idx, chart in enumerate(ws._charts):
    print(f"\n--- Chart {idx+1} ---")
    print(f"Title: {chart.title}")
    print(f"Type: {type(chart)}")
    print(f"Anchor: {chart.anchor}")
    if hasattr(chart, 'width'):
        print(f"Width: {chart.width}, Height: {chart.height}")

print("\n--- Rows Content ---")
for r in range(1, ws.max_row + 1):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    # If the row has any non-None value, print it
    if any(val is not None for val in row_vals):
        # Trim ending None values for readability
        while row_vals and row_vals[-1] is None:
            row_vals.pop()
        print(f"Row {r:03d}: {row_vals}")
