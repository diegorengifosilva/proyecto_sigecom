import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (63).xlsx", data_only=False)
ws = wb["Admin y Com"]
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
print("--- Sheet 'Admin y Com' Rows ---")
for r in range(1, ws.max_row + 1):
    vals = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        vals.append(v)
    # Print row if not completely empty
    if any(x is not None for x in vals):
        print(f"Row {r:02d}: {vals}")
