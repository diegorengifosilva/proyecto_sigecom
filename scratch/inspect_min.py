import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (10).xlsx", data_only=True)
ws = wb["MIN"]
print(f"--- Sheet: MIN (max_row={ws.max_row}) ---")
for r in range(1, ws.max_row + 1):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 19)]
    if any(row_vals):
        print(f"Row {r:02d}: {row_vals[:12]} | cols 13-18: {row_vals[12:18]}")
