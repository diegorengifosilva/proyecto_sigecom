import openpyxl
from openpyxl.chart.pie_chart import PieChart3D
from openpyxl.chart.reference import Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import Legend
from openpyxl.chart._3d import View3D
from openpyxl.chart.shapes import GraphicalProperties

wb = openpyxl.Workbook()
ws = wb.active

# Setup mock cells
ws.cell(row=39, column=3, value="Avance")
ws.cell(row=39, column=4, value="Faltante")
ws.cell(row=40, column=3, value=0.54)
ws.cell(row=40, column=4, value=0.46)

gauge_min = PieChart3D()
gauge_min.legend = Legend()
gauge_min.legend.position = "b" # Bottom legend

ref_vals_min = Reference(ws, min_col=3, max_col=4, min_row=40, max_row=40)
ref_cats = Reference(ws, min_col=3, max_col=4, min_row=39, max_row=39)

print("Adding data...")
gauge_min.add_data(ref_vals_min, from_rows=True)
gauge_min.set_categories(ref_cats)

print("Configuring labels...")
gauge_min.dataLabels = DataLabelList()
gauge_min.dataLabels.showVal = True
gauge_min.dataLabels.showPercent = False
gauge_min.dataLabels.showCatName = False # Disable category name to make it look clean like user's screenshot!
gauge_min.dataLabels.showSerName = False

# Layout setup
gauge_min.layout = Layout(
    manualLayout=ManualLayout(
        x=0.1, y=0.1,
        h=0.75, w=0.8,
        xMode="edge", yMode="edge"
    )
)

# 3D view
gauge_min.view3D = View3D(rotX=10, rotY=40, rAngAx=False)
gauge_min.view3D.perspective = 0
gauge_min.view3D.depthPercent = 130

# Colors
dp0 = DataPoint(idx=0)
dp0.graphicalProperties.solidFill = "237573"
dp1 = DataPoint(idx=1)
dp1.graphicalProperties.solidFill = "E2E8F0"

gauge_min.series[0].dPt.append(dp0)
gauge_min.series[0].dPt.append(dp1)

# Remove border
gauge_min.graphical_properties = GraphicalProperties()
gauge_min.graphical_properties.ln.noFill = True

wb.save("scratch/test_pie_out.xlsx")
print("Saved workbook successfully!")
