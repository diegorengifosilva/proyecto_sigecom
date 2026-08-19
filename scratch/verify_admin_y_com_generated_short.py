import openpyxl

wb = openpyxl.load_workbook("scratch/test_generated_report.xlsx", data_only=False)
ws = wb["Admin y Com"]

print("--- Admin y Com Formulas (Rows 1 to 21) ---")
for r in range(1, 22):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 12)]
    print(f"Row {r:02d} Formulas: {row_vals}")

print("\n--- Admin y Com Formulas (Rows 25 to 27) ---")
for r in range(25, 28):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 7)]
    print(f"Row {r:02d} Formulas: {row_vals}")

print("\n--- Admin y Com Formulas (Rows 50 to 54) ---")
for r in range(50, 55):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 7)]
    print(f"Row {r:02d} Formulas: {row_vals}")

wb_val = openpyxl.load_workbook("scratch/test_generated_report.xlsx", data_only=True)
ws_val = wb_val["Admin y Com"]

print("\n--- Admin y Com Values (Rows 1 to 21) ---")
for r in range(1, 22):
    row_vals = [ws_val.cell(row=r, column=c).value for c in range(1, 12)]
    print(f"Row {r:02d} Values: {row_vals}")
