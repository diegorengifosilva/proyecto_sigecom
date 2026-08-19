import openpyxl
import json

wb = openpyxl.load_workbook("SGC.REG-004 Seguimiento de OC 2026.xlsx", data_only=False)
ws = wb["Objetivos"]

# Print cells to understand the layout
print("\n--- Cells in 'Objetivos' sheet (Row 50 onwards, 10 columns) ---")
for r in range(50, 75):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 11)]
    if any(val is not None for val in row_vals):
        print(f"Row {r:02d}: {row_vals}")

# Export chart properties
print(f"\nTotal charts in 'Objetivos': {len(ws._charts)}")
charts_info = []
for idx, chart in enumerate(ws._charts):
    info = {
        "index": idx + 1,
        "class": chart.__class__.__name__,
        "width": getattr(chart, "width", None),
        "height": getattr(chart, "height", None),
        "anchor": None
    }
    
    # Extract anchor
    if hasattr(chart, 'anchor') and chart.anchor:
        try:
            if hasattr(chart.anchor, '_from') and chart.anchor._from:
                col = chart.anchor._from.col
                row = chart.anchor._from.row
                col_letter = openpyxl.utils.get_column_letter(col + 1)
                info["anchor"] = f"{col_letter}{row + 1}"
        except Exception:
            pass
            
    # Extract title
    title_text = None
    if chart.title:
        try:
            if hasattr(chart.title, 'text') and isinstance(chart.title.text, str):
                title_text = chart.title.text
            elif hasattr(chart.title, 'tx') and hasattr(chart.title.tx, 'rich') and chart.title.tx.rich:
                texts = []
                for p in chart.title.tx.rich.p:
                    for r in p.r:
                        if hasattr(r, 't') and isinstance(r.t, str):
                            texts.append(r.t)
                if texts:
                    title_text = "".join(texts)
        except Exception:
            pass
    info["title"] = title_text
    
    # Series details
    series_info = []
    for s_idx, s in enumerate(chart.series):
        t_ref = "None"
        if s.title:
            t_ref = s.title.text if hasattr(s.title, 'text') else str(s.title)
        s_val = {
            "index": s_idx + 1,
            "title": t_ref,
            "val_ref": s.val.numRef.f if hasattr(s, 'val') and s.val and hasattr(s.val, 'numRef') and s.val.numRef else 'None',
            "cat_ref": s.cat.strRef.f if hasattr(s, 'cat') and s.cat and hasattr(s.cat, 'strRef') and s.cat.strRef else 'None'
        }
        series_info.append(s_val)
    info["series"] = series_info
    
    charts_info.append(info)

with open("scratch/charts_sgc_objetivos.json", "w") as f:
    json.dump(charts_info, f, indent=2)
print("\nSuccessfully wrote scratch/charts_sgc_objetivos.json")
