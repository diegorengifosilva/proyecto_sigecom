import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (22).xlsx", data_only=False)
ws = wb["2012 - 2026"]
print(f"Sheet: {ws.title}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

print("\n--- Rows 140 to 180 Content ---")
for r in range(140, min(180, ws.max_row + 1)):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 10)]
    if any(val is not None for val in row_vals):
        while row_vals and row_vals[-1] is None:
            row_vals.pop()
        print(f"Row {r:03d}: {row_vals}")
