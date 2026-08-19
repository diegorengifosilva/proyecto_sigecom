import os
import sys
import django

# Add project root to sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'backend.settings')
django.setup()

import openpyxl
from django.test import RequestFactory
from dashboard_api.views import exportar_mensual

factory = RequestFactory()
request = factory.get('/api/dashboard/exportar_mensual/?anno=2026')

print("Generating Excel report...")
response = exportar_mensual(request)

if response.status_code == 200:
    out_filename = "scratch/test_generated_report.xlsx"
    with open(out_filename, "wb") as f:
        f.write(response.content)
    print(f"Report generated successfully and saved to {out_filename}")
    
    # Load and inspect the generated sheet
    wb = openpyxl.load_workbook(out_filename, data_only=False)
    ws = wb["2012 - 2026"]
    print(f"\n--- Sheet '2012 - 2026' ---")
    print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
    
    # Find GENERAL title row
    gen_row = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "GENERAL":
            gen_row = r
            break
            
    if gen_row:
        print(f"\nFound GENERAL table starting at row {gen_row}:")
        for r in range(gen_row, min(gen_row + 15, ws.max_row + 1)):
            row_vals = [ws.cell(row=r, column=c).value for c in range(1, 8)]
            while row_vals and row_vals[-1] is None:
                row_vals.pop()
            print(f"  Row {r:03d}: {row_vals}")
            
        print(f"\nCharts in worksheet: {len(ws._charts)}")
        for idx, chart in enumerate(ws._charts):
            title_text = "No Title"
            if chart.title:
                if hasattr(chart.title, 'text') and chart.title.text:
                    title_text = str(chart.title.text)
                elif hasattr(chart.title, 'tx') and hasattr(chart.title.tx, 'rich') and chart.title.tx.rich:
                    texts = []
                    for p in chart.title.tx.rich.p:
                        for r in p.r:
                            if hasattr(r, 't'):
                                texts.append(r.t)
                    title_text = "".join(texts)
                    
            anchor_cell = "Unknown"
            if hasattr(chart, 'anchor') and chart.anchor:
                try:
                    if hasattr(chart.anchor, '_from') and chart.anchor._from:
                        col_letter = openpyxl.utils.get_column_letter(chart.anchor._from.col + 1)
                        row_num = chart.anchor._from.row + 1
                        anchor_cell = f"{col_letter}{row_num}"
                except:
                    pass
            print(f"  Chart {idx+1:02d} | Title: {title_text.strip()} | Class: {chart.__class__.__name__} | Anchor: {anchor_cell} | Size: {chart.width}x{chart.height}")
            
        # Inspect Objetivos worksheet
        ws_obj = wb["Objetivos"]
        print(f"\n--- Sheet 'Objetivos' ---")
        print(f"Max row: {ws_obj.max_row}, Max col: {ws_obj.max_column}")
        for r in range(1, 21):
            row_vals = [ws_obj.cell(row=r, column=c).value for c in range(1, 14)]
            while row_vals and row_vals[-1] is None:
                row_vals.pop()
            if row_vals:
                print(f"  Row {r:02d}: {row_vals}")
                
        print(f"\nCharts in sheet 'Objetivos': {len(ws_obj._charts)}")
        for idx, chart in enumerate(ws_obj._charts):
            title_text = "No Title"
            if chart.title:
                if hasattr(chart.title, 'text') and chart.title.text:
                    title_text = str(chart.title.text)
                elif hasattr(chart.title, 'tx') and hasattr(chart.title.tx, 'rich') and chart.title.tx.rich:
                    texts = []
                    for p in chart.title.tx.rich.p:
                        for r in p.r:
                            if hasattr(r, 't'):
                                texts.append(r.t)
                    title_text = "".join(texts)
                    
            anchor_cell = "Unknown"
            if hasattr(chart, 'anchor') and chart.anchor:
                try:
                    if hasattr(chart.anchor, '_from') and chart.anchor._from:
                        col_letter = openpyxl.utils.get_column_letter(chart.anchor._from.col + 1)
                        row_num = chart.anchor._from.row + 1
                        anchor_cell = f"{col_letter}{row_num}"
                except:
                    pass
            print(f"  Chart {idx+1:02d} | Title: {title_text.strip()} | Class: {chart.__class__.__name__} | Anchor: {anchor_cell} | Size: {chart.width}x{chart.height}")
    else:
        print("ERROR: GENERAL table was not found in sheet '2012 - 2026'")
else:
    print(f"ERROR: View returned status code {response.status_code}")
    print(response.content)
