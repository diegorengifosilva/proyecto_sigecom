import openpyxl
from openpyxl.chart import BarChart3D

chart = BarChart3D()
print("Default width:", chart.width)
print("Default height:", chart.height)

chart.width = 15
chart.height = 9.54
print("Assigned width:", chart.width)
print("Assigned height:", chart.height)

wb = openpyxl.Workbook()
ws = wb.active
ws.add_chart(chart, "A1")
wb.save("scratch/test_dims_out.xlsx")

wb2 = openpyxl.load_workbook("scratch/test_dims_out.xlsx")
ws2 = wb2.active
chart2 = ws2._charts[0]
print("Loaded width:", chart2.width)
print("Loaded height:", chart2.height)
