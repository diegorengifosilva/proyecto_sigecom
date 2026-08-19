import openpyxl
from openpyxl.chart import BarChart3D, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.chart._3d import View3D

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Objetivos Com."
ws.views.sheetView[0].showGridLines = True

# Add dummy data starting at row 29
for _ in range(28):
    ws.append([])

ws.cell(row=29, column=2, value="Area")
ws.cell(row=29, column=3, value="Cuota Comercial")
ws.cell(row=29, column=4, value="Conseguido")
ws.cell(row=29, column=5, value="%")

ws.cell(row=30, column=2, value="Mineria")
ws.cell(row=30, column=3, value=625000)
ws.cell(row=30, column=4, value=435397.24)
ws.cell(row=30, column=5, value=0.696)

ws.cell(row=31, column=2, value="Industria")
ws.cell(row=31, column=3, value=260000)
ws.cell(row=31, column=4, value=14470.20)
ws.cell(row=31, column=5, value=0.056)

ws.cell(row=32, column=2, value="Petroquimica")
ws.cell(row=32, column=3, value=230000)
ws.cell(row=32, column=4, value=178769.53)
ws.cell(row=32, column=5, value=0.777)

ws.cell(row=33, column=2, value="Safety")
ws.cell(row=33, column=3, value=90000)
ws.cell(row=33, column=4, value=0)
ws.cell(row=33, column=5, value=0)

chart = BarChart3D()
chart.type = "col"
chart.grouping = "clustered"
chart.style = 10
chart.title = "Seguimiento Objetivo Comercial"

ref_data = Reference(ws, min_col=3, max_col=4, min_row=29, max_row=33)
cats = Reference(ws, min_col=2, min_row=30, max_row=33)

chart.add_data(ref_data, titles_from_data=True)
chart.set_categories(cats)

# Style series colors
if len(chart.series) > 0:
    chart.series[0].graphicalProperties.solidFill = "8FAADC"
if len(chart.series) > 1:
    chart.series[1].graphicalProperties.solidFill = "237573"

chart.legend = Legend()
chart.legend.position = "t"

chart.dataLabels = DataLabelList()
chart.dataLabels.showVal = True
chart.dataLabels.showCatName = False
chart.dataLabels.showSerName = False
chart.dataLabels.showPercent = False
chart.dataLabels.showLegendKey = False
chart.dataLabels.numFmt = '$ #,##0.00'

chart.y_axis.delete = True
chart.x_axis.delete = False
chart.y_axis.majorGridlines = None
chart.x_axis.majorGridlines = None

chart.view3D = View3D(rotX=10, rotY=40, rAngAx=False)
chart.view3D.perspective = 0
chart.view3D.depthPercent = 130

chart.width = 16.5
chart.height = 8.5

ws.add_chart(chart, "G29")
wb.save("scratch/test_chart_result.xlsx")
print("Saved scratch/test_chart_result.xlsx successfully!")
