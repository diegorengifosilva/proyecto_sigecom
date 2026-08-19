import openpyxl
from openpyxl.chart import BarChart, Reference

wb = openpyxl.Workbook()
ws = wb.active
ws.append(["Category", "Value"])
ws.append(["A", 10])
ws.append(["B", 20])

chart = BarChart()
chart.roundedCorners = True
data = Reference(ws, min_col=2, min_row=1, max_row=3)
cats = Reference(ws, min_col=1, min_row=2, max_row=3)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws.add_chart(chart, "D2")

wb.save("scratch/test_rounded_chart.xlsx")

# Load it back
wb2 = openpyxl.load_workbook("scratch/test_rounded_chart.xlsx")
ws2 = wb2.active
ch = ws2._charts[0]
print("Loaded chart roundedCorners:", ch.roundedCorners)

# Read the XML file inside the zip to see if roundedCorners is present!
import zipfile
with zipfile.ZipFile("scratch/test_rounded_chart.xlsx", "r") as z:
    for filename in z.namelist():
        if "charts/chart" in filename:
            xml_content = z.read(filename).decode("utf-8")
            print(f"--- XML of {filename} ---")
            print("roundedCorners in XML:", "roundedCorners" in xml_content)
            # Find the position of roundedCorners in XML
            idx = xml_content.find("roundedCorners")
            if idx != -1:
                print("Snippet:", xml_content[idx-20:idx+40])
