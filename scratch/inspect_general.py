import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (22).xlsx", data_only=False)
ws = wb["2012 - 2026"]
print(f"Sheet: {ws.title}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
print(f"Number of charts: {len(ws._charts)}")

print("\n--- Rows 140 to 175 Content ---")
for r in range(140, min(175, ws.max_row + 1)):
    row_vals = [ws.cell(row=r, column=c).value for c in range(1, 15)]
    if any(val is not None for val in row_vals):
        while row_vals and row_vals[-1] is None:
            row_vals.pop()
        print(f"Row {r:03d}: {row_vals}")

print("\n--- Charts detail ---")
for idx, chart in enumerate(ws._charts):
    title_text = ""
    if chart.title:
        # handle Text / Title objects
        title_text = str(chart.title.text if hasattr(chart.title, 'text') else chart.title)
    
    print(f"\nChart {idx+1}:")
    print(f"  Title: {title_text}")
    print(f"  Type: {chart.__class__.__name__}")
    
    # Try to find anchor coordinates
    anchor_str = "Unknown"
    if hasattr(chart, 'anchor') and chart.anchor:
        if hasattr(chart.anchor, '_from') and chart.anchor._from:
            col_letter = openpyxl.utils.get_column_letter(chart.anchor._from.col + 1)
            row_num = chart.anchor._from.row + 1
            anchor_str = f"{col_letter}{row_num}"
        else:
            anchor_str = str(chart.anchor)
    print(f"  Anchor: {anchor_str}")
    
    if hasattr(chart, 'width'):
        print(f"  Width: {chart.width}, Height: {chart.height}")
    
    if hasattr(chart, 'series') and chart.series:
        print(f"  Number of series: {len(chart.series)}")
        for s_idx, s in enumerate(chart.series):
            # values reference
            val_ref = "None"
            if hasattr(s, 'values') and s.values:
                if hasattr(s.values, 'numRef') and s.values.numRef:
                    val_ref = s.values.numRef.f
                elif hasattr(s.values, 'strRef') and s.values.strRef:
                    val_ref = s.values.strRef.f
            
            # categories/xVal reference
            cat_ref = "None"
            if hasattr(chart, 'x_axis') and hasattr(s, 'categories') and s.categories:
                if hasattr(s.categories, 'numRef') and s.categories.numRef:
                    cat_ref = s.categories.numRef.f
                elif hasattr(s.categories, 'strRef') and s.categories.strRef:
                    cat_ref = s.categories.strRef.f
                    
            title_ref = "None"
            if hasattr(s, 'title') and s.title:
                if hasattr(s.title, 'strRef') and s.title.strRef:
                    title_ref = s.title.strRef.f
                else:
                    title_ref = str(s.title)
                    
            print(f"    Series {s_idx}: title_ref={title_ref}, values_ref={val_ref}, categories_ref={cat_ref}")
