import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (10).xlsx")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n===== Sheet: {name} =====")
    for r in [2, 3, 4]:
        vals = [ws.cell(row=r, column=c).value for c in range(12, 19)]
        print(f"Row {r}: cols 12-18={vals}")
