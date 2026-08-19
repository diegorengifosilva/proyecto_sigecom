import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (63).xlsx", data_only=False)
if "Admin y Com" in wb.sheetnames:
    ws = wb["Admin y Com"]
    print("--- Admin y Com Sheet in reporte_mensual (63).xlsx (Formulas) ---")
    print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
    print(f"Charts in sheet: {len(ws._charts)}")
    
    # Print the first 100 rows and 15 columns
    for r in range(1, min(ws.max_row + 1, 101)):
        row_vals = []
        for c in range(1, min(ws.max_column + 1, 16)):
            val = ws.cell(row=r, column=c).value
            if val is not None:
                row_vals.append((c, val))
        if row_vals:
            print(f"Row {r:02d}: {row_vals}")

    # Let's inspect data_only values too
    wb_val = openpyxl.load_workbook("reporte_mensual (63).xlsx", data_only=True)
    ws_val = wb_val["Admin y Com"]
    print("\n--- Admin y Com Sheet in reporte_mensual (63).xlsx (Values) ---")
    for r in range(1, min(ws_val.max_row + 1, 101)):
        row_vals = []
        for c in range(1, min(ws_val.max_column + 1, 16)):
            val = ws_val.cell(row=r, column=c).value
            if val is not None:
                row_vals.append((c, val))
        if row_vals:
            print(f"Row {r:02d}: {row_vals}")
            
    # Print details of any charts present
    if len(ws._charts) > 0:
        for idx, chart in enumerate(ws._charts):
            print(f"\nChart {idx}: {type(chart)}")
            print("  Title:", getattr(chart.title, "text", "No Title") if chart.title else "No Title")
            print("  Anchor:", chart.anchor)
            print("  Width:", chart.width)
            print("  Height:", chart.height)
            print("  Series Count:", len(chart.series))
            for s_idx, s in enumerate(chart.series):
                print(f"    Series {s_idx} Title: {s.title}")
                if hasattr(s, "val") and s.val:
                    print(f"      Val Ref:", s.val.numRef.f if getattr(s.val, "numRef", None) else "No numRef")
else:
    print("Admin y Com sheet not found in reporte_mensual (63).xlsx")
