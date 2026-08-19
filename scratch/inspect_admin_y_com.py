import openpyxl

wb = openpyxl.load_workbook("SGC.REG-004 Seguimiento de OC 2026.xlsx", data_only=False)
if "Admin y Com" in wb.sheetnames:
    ws = wb["Admin y Com"]
    print("--- Admin y Com Sheet Inspection (Non-Evaluated Formulas) ---")
    print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
    print(f"Charts in sheet: {len(ws._charts)}")
    
    # Print the first 100 rows and 15 columns
    for r in range(1, min(ws.max_row + 1, 101)):
        row_vals = []
        for c in range(1, min(ws.max_column + 1, 16)):
            val = ws.cell(row=r, column=c).value
            if val is not None:
                row_vals.append((c, val))
        if row_vals:
            print(f"Row {r:02d}: {row_vals}")

    # Let's also open with data_only=True to see the values
    wb_val = openpyxl.load_workbook("SGC.REG-004 Seguimiento de OC 2026.xlsx", data_only=True)
    ws_val = wb_val["Admin y Com"]
    print("\n--- Admin y Com Sheet Inspection (Evaluated Values) ---")
    for r in range(1, min(ws_val.max_row + 1, 101)):
        row_vals = []
        for c in range(1, min(ws_val.max_column + 1, 16)):
            val = ws_val.cell(row=r, column=c).value
            if val is not None:
                row_vals.append((c, val))
        if row_vals:
            print(f"Row {r:02d}: {row_vals}")
else:
    print("Admin y Com sheet not found in SGC.REG-004 Seguimiento de OC 2026.xlsx")
