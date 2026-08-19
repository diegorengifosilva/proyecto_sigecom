import openpyxl
from openpyxl.chart.pie_chart import DoughnutChart
from openpyxl.chart.reference import Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout

wb = openpyxl.Workbook()
ws = wb.active

# Setup mock cells
ws.cell(row=39, column=3, value="Avance")
ws.cell(row=39, column=4, value="Faltante")
ws.cell(row=40, column=3, value=0.54)
ws.cell(row=40, column=4, value=0.46)

gauge_min = DoughnutChart()
gauge_min.holeSize = 70
gauge_min.legend = None

ref_vals_min = Reference(ws, min_col=3, max_col=4, min_row=40, max_row=40)
ref_cats = Reference(ws, min_col=3, max_col=4, min_row=39, max_row=39)

print("Adding data...")
gauge_min.add_data(ref_vals_min, from_rows=True)
print("Setting categories...")
gauge_min.set_categories(ref_cats)

print("Configuring labels...")
gauge_min.dataLabels = DataLabelList()
gauge_min.dataLabels.showVal = True
gauge_min.dataLabels.showPercent = False
gauge_min.dataLabels.showCatName = True
gauge_min.dataLabels.showSerName = False

# Layout setup
gauge_min.layout = Layout(
    manualLayout=ManualLayout(
        x=0.1, y=0.15,
        h=0.8, w=0.8,
        xMode="edge", yMode="edge"
    )
)

dp0 = DataPoint(idx=0)
dp0.graphicalProperties.solidFill = "237573"
dp1 = DataPoint(idx=1)
dp1.graphicalProperties.solidFill = "E2E8F0"

gauge_min.series[0].dPt.append(dp0)
gauge_min.series[0].dPt.append(dp1)

wb.save("scratch/test_out.xlsx")
print("Saved workbook successfully!")
