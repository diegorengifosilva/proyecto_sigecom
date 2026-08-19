import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (50).xlsx", data_only=False)
print("Worksheet names:", wb.sheetnames)

if "Objetivos" in wb.sheetnames:
    ws = wb["Objetivos"]
    print("\n--- Cells in 'Objetivos' sheet (Row 45 onwards) ---")
    for r in range(45, 80):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 15)]
        if any(val is not None for val in row_vals):
            print(f"Row {r:02d}: {row_vals}")
            
    print(f"\nTotal charts in 'Objetivos' of report 50: {len(ws._charts)}")
    for idx, chart in enumerate(ws._charts):
        anchor_val = None
        if hasattr(chart, 'anchor') and chart.anchor:
            try:
                if hasattr(chart.anchor, '_from') and chart.anchor._from:
                    col = chart.anchor._from.col
                    row = chart.anchor._from.row
                    col_letter = openpyxl.utils.get_column_letter(col + 1)
                    anchor_val = f"{col_letter}{row + 1}"
            except Exception:
                pass
        title_text = chart.title.text if hasattr(chart, "title") and chart.title else "None"
        print(f"  Chart {idx+1}: class={chart.__class__.__name__}, anchor={anchor_val}, title={title_text}")
