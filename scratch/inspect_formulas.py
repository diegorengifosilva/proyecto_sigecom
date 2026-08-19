import openpyxl

wb = openpyxl.load_workbook("SGC.REG-004 Seguimiento de OC 2026.xlsx", data_only=False)
ws = wb["Objetivos"]

print("--- Formulas in SGC.REG-004 Seguimiento de OC 2026.xlsx ---")
for r in range(62, 73):
    row_vals = []
    for c in range(1, 10):
        val = ws.cell(row=r, column=c).value
        row_vals.append(val)
    print(f"Row {r:02d}: {row_vals}")

wb_val = openpyxl.load_workbook("SGC.REG-004 Seguimiento de OC 2026.xlsx", data_only=True)
ws_val = wb_val["Objetivos"]

print("\n--- Calculated Values in SGC.REG-004 Seguimiento de OC 2026.xlsx ---")
for r in range(62, 73):
    row_vals = []
    for c in range(1, 10):
        val = ws_val.cell(row=r, column=c).value
        row_vals.append(val)
    print(f"Row {r:02d}: {row_vals}")
