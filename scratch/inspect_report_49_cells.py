import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (49).xlsx", data_only=False)
ws = wb["Objetivos"]

with open("scratch/cells_49.txt", "w") as f:
    f.write("=== Cells (Formulas) ===\n")
    for r in range(45, 80):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 15)]
        f.write(f"Row {r:02d}: {row_vals}\n")

wb_val = openpyxl.load_workbook("reporte_mensual (49).xlsx", data_only=True)
ws_val = wb_val["Objetivos"]
with open("scratch/cells_49_val.txt", "w") as f:
    f.write("=== Cells (Calculated Values) ===\n")
    for r in range(45, 80):
        row_vals = [ws_val.cell(row=r, column=c).value for c in range(1, 15)]
        f.write(f"Row {r:02d}: {row_vals}\n")

print("Done writing scratch/cells_49.txt and scratch/cells_49_val.txt")
