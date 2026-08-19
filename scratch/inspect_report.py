import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (22).xlsx")
print("Sheets in workbook:", wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\nSheet: {sheetname}")
    print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
    print(f"Charts in sheet: {len(ws._charts)}")
    # Print the first few rows to see what is there
    for r in range(1, min(40, ws.max_row + 1)):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, min(15, ws.max_column + 1))]
        if any(row_vals):
            print(f"Row {r:02d}: {row_vals}")
