import openpyxl

wb = openpyxl.load_workbook("reporte_mensual (50).xlsx", data_only=False)
ws = wb["Objetivos"]

with open("scratch/cells_50.txt", "w") as f:
    f.write("=== Cells (Formulas) ===\n")
    for r in range(35, 80):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 15)]
        f.write(f"Row {r:02d}: {row_vals}\n")

wb_val = openpyxl.load_workbook("reporte_mensual (50).xlsx", data_only=True)
ws_val = wb_val["Objetivos"]
with open("scratch/cells_50_val.txt", "w") as f:
    f.write("=== Cells (Calculated Values) ===\n")
    for r in range(35, 80):
        row_vals = [ws_val.cell(row=r, column=c).value for c in range(1, 15)]
        f.write(f"Row {r:02d}: {row_vals}\n")

print("Done writing scratch/cells_50.txt and scratch/cells_50_val.txt")
