import openpyxl
from openpyxl.chart.pie_chart import DoughnutChart
from openpyxl.chart.scatter_chart import ScatterChart
from openpyxl.chart.series import XYSeries, Series
from openpyxl.chart.series import DataPoint
from openpyxl.chart.reference import Reference

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Objetivos"

# Populate dummy data
ws.cell(row=51, column=2, value="Porcentaje")
ws.cell(row=51, column=3, value="Escala")
for i in range(1, 10):
    ws.cell(row=51+i, column=2, value=i*0.1)
    ws.cell(row=51+i, column=3, value=1)
ws.cell(row=61, column=2, value=1.0)
ws.cell(row=61, column=3, value=9)

ws.cell(row=66, column=2, value="Inicio")
ws.cell(row=66, column=3, value=0)
ws.cell(row=66, column=4, value=0)
ws.cell(row=67, column=2, value="Fin")
ws.cell(row=67, column=3, value=0.5)
ws.cell(row=67, column=4, value=0.8)

# 1. Doughnut
chart = DoughnutChart()
chart.firstSliceAng = 270
chart.holeSize = 75

ref_vals = Reference(ws, min_col=3, min_row=52, max_row=61)
chart.add_data(ref_vals)

# Setup segment colors
series = chart.series[0]
for idx in range(10):
    dp = DataPoint(idx=idx)
    if idx == 9:
        dp.graphicalProperties.noFill = True
    else:
        dp.graphicalProperties.solidFill = "FF0000" if idx < 3 else ("FFFF00" if idx < 7 else "92D050")
    series.dPt.append(dp)

# 2. Scatter (needle)
scatter = ScatterChart()
scatter.x_axis.delete = True
scatter.y_axis.delete = True

xvals = Reference(ws, min_col=3, min_row=66, max_row=67)
yvals = Reference(ws, min_col=4, min_row=66, max_row=67)

needle_series = XYSeries()
needle_series.xValues = xvals
needle_series.yValues = yvals
needle_series.graphicalProperties.ln.w = 31750
needle_series.graphicalProperties.ln.solidFill = "002060"
needle_series.marker.symbol = "circle"
needle_series.marker.size = 5

scatter.series.append(needle_series)

# Combine
chart += scatter

ws.add_chart(chart, "E51")
wb.save("scratch/test_gauge.xlsx")
print("Successfully generated test_gauge.xlsx")
