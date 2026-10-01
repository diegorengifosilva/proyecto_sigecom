"""
Views para el sistema de cronogramas EDT/Gantt
Incluye ViewSets para versiones de cronogramas, tareas y documentos
"""

from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.parsers import MultiPartParser, FormParser
from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from datetime import datetime

from .models import CronogramaVersion, Tarea, DocumentoProyecto, ProyectoEv
from .serializers import CronogramaVersionSerializer, TareaSerializer, DocumentoProyectoSerializer
from .services_cronograma import compute_schedule, detect_cycles


class CronogramaVersionViewSet(viewsets.ModelViewSet):
    """
    ViewSet para versiones de cronogramas.

    Endpoints:
    - GET /api/proyectos/{id}/cronogramas/ - Listar versiones
    - POST /api/proyectos/{id}/cronogramas/ - Crear nueva versión
    - GET /api/proyectos/{id}/cronogramas/{v_id}/ - Detalle de versión
    - PUT /api/proyectos/{id}/cronogramas/{v_id}/ - Actualizar versión
    - DELETE /api/proyectos/{id}/cronogramas/{v_id}/ - Eliminar versión
    - POST /api/proyectos/{id}/cronogramas/{v_id}/calcular/ - Recalcular cronograma
    - GET /api/proyectos/{id}/cronogramas/{v_id}/export_msp/ - Exportar MS Project
    - GET /api/proyectos/{id}/cronogramas/{v_id}/export_xlsx/ - Exportar Excel
    """

    serializer_class = CronogramaVersionSerializer
    permission_classes = [AllowAny]
    pagination_class = None

    def get_queryset(self):
        proyecto_id = self.kwargs.get('proyecto_pk') or self.request.query_params.get('proyecto_id') or self.request.query_params.get('proyecto')
        if proyecto_id:
            return CronogramaVersion.objects.filter(
                proyecto_id=proyecto_id
            ).prefetch_related('tareas')
        return CronogramaVersion.objects.all()

    def perform_create(self, serializer):
        proyecto_id = self.kwargs.get('proyecto_pk') or self.request.data.get('proyecto')
        proyecto = get_object_or_404(ProyectoEv, pk=proyecto_id, usuario=self.request.user)
        serializer.save(proyecto=proyecto, creado_por=self.request.user)

    @action(detail=True, methods=['post'])
    def calcular(self, request, pk=None, proyecto_pk=None):
        """
        POST /api/proyectos/{id}/cronogramas/{v_id}/calcular/

        Recalcula el cronograma completo usando el algoritmo de programación.
        Valida que no haya dependencias circulares antes de calcular.
        """
        version = self.get_object()

        # Validar ciclos
        tareas = list(version.tareas.all())
        if detect_cycles(tareas):
            return Response(
                {"error": "Dependencias circulares detectadas en el cronograma"},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Calcular cronograma
        project_start = version.proyecto.fecha_inicio
        if not project_start:
            project_start = datetime.now()
        else:
            project_start = datetime.combine(project_start, datetime.min.time())

        tareas_calculadas = compute_schedule(version.tareas.all(), project_start)

        # Guardar cambios
        for tarea in tareas_calculadas:
            tarea.save()

        # Retornar versión actualizada
        serializer = self.get_serializer(version)
        return Response(serializer.data)

    @action(detail=True, methods=['get'])
    def export_msp(self, request, pk=None, proyecto_pk=None):
        """
        GET /api/proyectos/{id}/cronogramas/{v_id}/export_msp/

        Exporta el cronograma a formato MS Project XML.
        """
        from .exporters import export_to_msp_xml

        version = self.get_object()
        xml_content = export_to_msp_xml(version)

        response = HttpResponse(xml_content, content_type='application/xml')
        filename = f"{version.nombre.replace(' ', '_')}.xml"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        return response

    @action(detail=False, methods=['get'])
    def plantilla_excel(self, request, proyecto_pk=None):
        """
        GET /api/proyectos/{id}/cronogramas/plantilla_excel/

        Descarga una plantilla Excel vacía para importar cronograma.
        Incluye headers y una fila de ejemplo con instrucciones.
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from io import BytesIO

        # Crear workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Plantilla EDT"

        # Headers con formato
        headers = [
            'Item', 'EDT', 'Nivel', 'Nombre', 'Duración',
            'Calendario', 'Predecesoras', 'Hito', 'Timeline', 'Mostrar'
        ]

        # Agregar headers con estilo
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Fila de ejemplo con instrucciones
        ejemplo = [
            1,
            "1.0",
            1,
            "Ejemplo de Tarea",
            "5d",
            "8+1",
            "",
            "No",
            "No",
            "Si"
        ]
        ws.append(ejemplo)

        # Agregar instrucciones en una hoja separada
        ws_instrucciones = wb.create_sheet("Instrucciones")
        instrucciones = [
            ["INSTRUCCIONES PARA IMPORTAR CRONOGRAMA"],
            [""],
            ["Columna", "Descripción", "Formato", "Ejemplo"],
            ["Item", "Número consecutivo de la tarea", "Entero", "1, 2, 3..."],
            ["EDT", "Código de estructura de desglose", "Texto", "1.0, 1.1, 1.2.1"],
            ["Nivel", "Nivel jerárquico (1-4)", "Entero", "1, 2, 3, 4"],
            ["Nombre", "Nombre de la tarea", "Texto", "Diseño de planos"],
            ["Duración", "Duración de la tarea", "Texto", "5d, 16h, 2w"],
            ["", "  d = días, h = horas, w = semanas", "", ""],
            ["Calendario", "Tipo de calendario laboral", "Texto", "8+1, 10, 12"],
            ["Predecesoras", "Dependencias (opcional)", "Texto", "1:FS, 2:SS+2d"],
            ["", "  FS = Fin-Inicio", "", ""],
            ["", "  SS = Inicio-Inicio", "", ""],
            ["", "  FF = Fin-Fin", "", ""],
            ["", "  SF = Inicio-Fin", "", ""],
            ["Hito", "¿Es un hito?", "Si/No", "Si, No"],
            ["Timeline", "¿Mostrar en timeline?", "Si/No", "Si, No"],
            ["Mostrar", "¿Mostrar en Gantt?", "Si/No", "Si, No"],
            [""],
            ["NOTAS IMPORTANTES:"],
            ["- No modificar los headers de la primera fila"],
            ["- Cada tarea debe tener un Item único"],
            ["- El EDT debe seguir estructura jerárquica (1.0, 1.1, 1.2)"],
            ["- Las predecesoras se refieren al Item, no al EDT"],
            ["- Los hitos tienen duración 0 automáticamente"],
        ]

        for row in instrucciones:
            ws_instrucciones.append(row)

        # Ajustar anchos de columnas
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        # Guardar en buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        # Respuesta
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="Plantilla_Cronograma.xlsx"'

        return response

    @action(detail=False, methods=['post'], parser_classes=[MultiPartParser, FormParser])
    def importar_excel(self, request, proyecto_pk=None):
        """
        POST /api/proyectos/{id}/cronogramas/importar_excel/

        Importa un cronograma desde un archivo Excel.
        Crea una nueva versión del cronograma con las tareas del archivo.
        """
        import openpyxl
        from .services_cronograma import parse_duration_to_minutes, compute_schedule

        # Validar que se envió un archivo
        if 'file' not in request.FILES:
            return Response(
                {"error": "No se encontró el archivo. Usa el campo 'file' en la petición."},
                status=status.HTTP_400_BAD_REQUEST
            )

        archivo = request.FILES['file']

        # Validar extensión
        if not archivo.name.endswith(('.xlsx', '.xls')):
            return Response(
                {"error": "El archivo debe ser Excel (.xlsx o .xls)"},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            # Obtener proyecto
            proyecto = get_object_or_404(ProyectoEv, pk=proyecto_pk, usuario=request.user)

            # Leer archivo Excel
            wb = openpyxl.load_workbook(archivo)
            ws = wb.active

            # Leer headers (primera fila)
            headers = [cell.value for cell in ws[1]]

            # Mapear headers a índices
            col_map = {header: idx for idx, header in enumerate(headers)}

            # Validar headers requeridos
            required = ['Item', 'EDT', 'Nivel', 'Nombre', 'Duración', 'Calendario']
            missing = [h for h in required if h not in col_map]
            if missing:
                return Response(
                    {"error": f"Faltan columnas requeridas: {', '.join(missing)}"},
                    status=status.HTTP_400_BAD_REQUEST
                )

            # Crear nueva versión de cronograma
            nombre_version = request.data.get('nombre', f"Importado {datetime.now().strftime('%Y-%m-%d %H:%M')}")
            version = CronogramaVersion.objects.create(
                proyecto=proyecto,
                nombre=nombre_version,
                creado_por=request.user,
                es_activa=True
            )

            # Leer tareas (desde fila 2 en adelante)
            tareas_creadas = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[col_map['Item']]:  # Saltar filas vacías
                    continue

                # Parsear valores
                item = int(row[col_map['Item']])
                edt = str(row[col_map['EDT']]).strip()
                nivel = int(row[col_map['Nivel']])
                nombre = str(row[col_map['Nombre']]).strip()
                duracion_str = str(row[col_map['Duración']]).strip()
                calendario = str(row[col_map['Calendario']]).strip()

                # Valores opcionales
                predecesoras = str(row[col_map.get('Predecesoras', len(headers))] or '').strip()
                hito = str(row[col_map.get('Hito', len(headers))] or 'No').lower() in ['si', 'sí', 'yes', 'true', '1']
                timeline = str(row[col_map.get('Timeline', len(headers))] or 'No').lower() in ['si', 'sí', 'yes', 'true', '1']
                mostrar = str(row[col_map.get('Mostrar', len(headers))] or 'Si').lower() in ['si', 'sí', 'yes', 'true', '1']

                # Convertir duración a minutos
                duracion_min = parse_duration_to_minutes(duracion_str)

                # Crear tarea
                tarea = Tarea.objects.create(
                    version=version,
                    item=item,
                    edt=edt,
                    nivel=nivel,
                    nombre=nombre,
                    duracion_min=duracion_min,
                    calendario=calendario,
                    predecesoras=predecesoras,
                    es_hito=hito,
                    timeline=timeline,
                    mostrar=mostrar,
                    es_resumen=(nivel == 1)  # Tareas de nivel 1 son resumen
                )
                tareas_creadas.append(tarea)

            # Calcular cronograma automáticamente
            project_start = proyecto.fecha_inicio
            if not project_start:
                project_start = datetime.now()
            else:
                project_start = datetime.combine(project_start, datetime.min.time())

            tareas_calculadas = compute_schedule(version.tareas.all(), project_start)

            # Guardar cambios
            for tarea in tareas_calculadas:
                tarea.save()

            # Retornar versión creada
            serializer = self.get_serializer(version)
            return Response({
                "message": f"Cronograma importado exitosamente con {len(tareas_creadas)} tareas",
                "cronograma": serializer.data
            }, status=status.HTTP_201_CREATED)

        except openpyxl.utils.exceptions.InvalidFileException:
            return Response(
                {"error": "El archivo Excel está corrupto o no es válido"},
                status=status.HTTP_400_BAD_REQUEST
            )
        except Exception as e:
            return Response(
                {"error": f"Error al importar cronograma: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @action(detail=True, methods=['get'])
    def export_xlsx(self, request, pk=None, proyecto_pk=None):
        """
        GET /api/proyectos/{id}/cronogramas/{v_id}/export_xlsx/

        Exporta el cronograma a formato Excel.
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from io import BytesIO

        version = self.get_object()

        # Crear workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "EDT"

        # Headers con formato
        headers = [
            'Item', 'EDT', 'Nivel', 'Nombre de Tarea', 'Comienzo',
            'Fin', 'Duración (min)', 'Calendario', 'Hito', 'Predecesoras',
            'Timeline', 'Mostrar', 'Crítica', 'Resumen'
        ]

        # Aplicar estilo a headers
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

        # Data
        for tarea in version.tareas.all().order_by('item'):
            ws.append([
                tarea.item,
                tarea.edt,
                tarea.nivel,
                tarea.nombre,
                tarea.inicio.strftime('%Y-%m-%d %H:%M') if tarea.inicio else '',
                tarea.fin.strftime('%Y-%m-%d %H:%M') if tarea.fin else '',
                tarea.duracion_min,
                tarea.calendario,
                'Sí' if tarea.es_hito else 'No',
                tarea.predecesoras,
                'Sí' if tarea.timeline else 'No',
                'Sí' if tarea.mostrar else 'No',
                'Sí' if tarea.es_critica else 'No',
                'Sí' if tarea.es_resumen else 'No',
            ])

        # Ajustar anchos
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        # Guardar en buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        # Respuesta
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"{version.nombre.replace(' ', '_')}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        return response


class TareaViewSet(viewsets.ModelViewSet):
    """
    ViewSet para tareas de cronograma.
    """

    serializer_class = TareaSerializer
    permission_classes = [AllowAny]
    pagination_class = None

    def get_queryset(self):
        version_id = self.request.query_params.get('version_id')
        if version_id:
            return Tarea.objects.filter(version_id=version_id)
        return Tarea.objects.all()


class DocumentoProyectoViewSet(viewsets.ModelViewSet):
    """
    ViewSet para documentos/archivos del proyecto.

    Soporta upload de archivos mediante multipart/form-data.
    """

    serializer_class = DocumentoProyectoSerializer
    permission_classes = [AllowAny]
    pagination_class = None
    parser_classes = [MultiPartParser, FormParser]

    def get_queryset(self):
        proyecto_id = self.request.query_params.get('proyecto_id') or self.request.query_params.get('proyecto')
        if proyecto_id:
            return DocumentoProyecto.objects.filter(
                proyecto_id=proyecto_id
            ).select_related('subido_por')
        return DocumentoProyecto.objects.all()

    def perform_create(self, serializer):
        archivo = self.request.FILES.get('archivo')

        # Calcular tamaño
        tamano = archivo.size if archivo else 0

        # Detectar tipo por extensión
        nombre = archivo.name if archivo else ''
        tipo = 'documento'

        ext = nombre.lower().split('.')[-1] if '.' in nombre else ''

        if ext in ('jpg', 'jpeg', 'png', 'gif', 'bmp', 'svg', 'webp'):
            tipo = 'imagen'
        elif ext in ('pdf',):
            tipo = 'pdf'
        elif ext in ('xls', 'xlsx', 'xlsm'):
            tipo = 'excel'
        elif ext in ('dwg', 'dxf', 'pdf'):
            tipo = 'plano'

        serializer.save(
            subido_por=self.request.user,
            tamano_bytes=tamano,
            tipo=tipo
        )
