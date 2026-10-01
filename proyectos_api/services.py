# services.py
import openpyxl
from .models import TareaEv, HistorialEv, CurvaEv, ProyectoEv
from datetime import timedelta
from django.utils import timezone
from django.utils.timezone import make_aware
from django.db import transaction

def calcular_curva_s(proyecto):
    tareas = TareaEv.objects.filter(proyecto=proyecto)
    historiales = HistorialEv.objects.filter(proyecto=proyecto)

    fechas_set = set()
    for t in tareas:
        if getattr(t, 'fecha_inicio', None):
            fechas_set.add(t.fecha_inicio)
    for h in historiales:
        if getattr(h, 'fecha_registro', None):
            fechas_set.add(h.fecha_registro)

    if not fechas_set:
        inicio_p = getattr(proyecto, 'fecha_inicio', None) or timezone.now().date()
        fechas_set.add(inicio_p)

    fechas = sorted(fechas_set)
    resultados = []

    for f in fechas:
        pv = sum(getattr(t, 'valor_planificado', 0) or 0 for t in tareas if getattr(t, 'fecha_inicio', f) <= f)
        ev = sum(getattr(h, 'valor_ganado', 0) or 0 for h in historiales if getattr(h, 'fecha_registro', f) <= f)
        ac = sum(getattr(h, 'costo_real', 0) or 0 for h in historiales if getattr(h, 'fecha_registro', f) <= f)

        curva, _ = CurvaEv.objects.update_or_create(
            proyecto=proyecto,
            fecha=f,
            defaults={'pv': pv, 'ev': ev, 'ac': ac, 'fuente': 'calculado'}
        )
        resultados.append(curva)

    return resultados

def actualizar_metricas_ev(proyecto):
    """
    Actualiza automáticamente las métricas EVM (EV, PV, AC, SPI, CPI)
    en el modelo ProyectoEv a partir de las curvas existentes.
    """
    curva_reciente = CurvaEv.objects.filter(proyecto=proyecto).order_by('-fecha').first()
    
    if curva_reciente:
        proyecto.valor_ganado = curva_reciente.ev or 0
        proyecto.valor_planificado = curva_reciente.pv or 0
        proyecto.costo_real = curva_reciente.ac or 0
    else:
        proyecto.valor_ganado = proyecto.valor_ganado or 0
        proyecto.valor_planificado = proyecto.valor_planificado or 0
        proyecto.costo_real = proyecto.costo_real or 0

    # Calcular SPI y CPI si es posible
    if proyecto.valor_planificado and proyecto.valor_planificado > 0:
        proyecto.spi = proyecto.valor_ganado / proyecto.valor_planificado
    else:
        proyecto.spi = None

    if proyecto.costo_real and proyecto.costo_real > 0:
        proyecto.cpi = proyecto.valor_ganado / proyecto.costo_real
    else:
        proyecto.cpi = None

    proyecto.save()
    return proyecto

def generar_cronograma(proyecto):
    tareas = TareaEv.objects.filter(proyecto=proyecto).order_by('fecha_inicio')

    cronograma = []
    for tarea in tareas:
        duracion = None
        if tarea.fecha_inicio and tarea.fecha_fin:
            duracion = (tarea.fecha_fin - tarea.fecha_inicio).days

        cronograma.append({
            'id': tarea.id,
            'nombre': tarea.nombre,
            'fecha_inicio': tarea.fecha_inicio,
            'fecha_fin': tarea.fecha_fin,
            'duracion_dias': duracion,
            'responsable': tarea.responsable,
            'estado': tarea.estado,
            'descripcion': tarea.descripcion,
            'tipo': tarea.tipo,
        })

    return cronograma

def cargar_cronograma_desde_excel(archivo_excel, proyecto_id):
    wb = openpyxl.load_workbook(archivo_excel)
    hoja = wb.active
    proyecto = ProyectoEv.objects.get(id=proyecto_id)

    tareas_dict = {}

    with transaction.atomic():
        for i, row in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
            codigo = row[0]
            nombre = row[1]
            duracion = row[3]
            predecesoras = row[4]  # Puede ser '1,2,3' o None

            if not nombre:
                continue

            # Evitar duplicados por nombre dentro del mismo proyecto
            if TareaEv.objects.filter(nombre=nombre, proyecto=proyecto).exists():
                continue

            tarea = TareaEv.objects.create(
                proyecto=proyecto,
                nombre=nombre,
                duracion=int(duracion or 0),
                fecha_inicio=None,
                fecha_fin=None,
                es_predefinida=False  # El usuario puede editar luego
            )
            tareas_dict[codigo] = tarea

        # Segunda pasada para asignar predecesoras y calcular fechas
        for i, row in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
            codigo = row[0]
            predecesoras = row[4]

            tarea = tareas_dict.get(codigo)
            if not tarea:
                continue

            if predecesoras:
                lista = [t.strip() for t in str(predecesoras).split(",")]
                for cod in lista:
                    pred = tareas_dict.get(cod)
                    if pred:
                        tarea.predecesoras.add(pred)

            # Calcular fecha_inicio y fecha_fin por predecesoras
            if tarea.predecesoras.exists():
                fecha_inicio = max([p.fecha_fin for p in tarea.predecesoras.all() if p.fecha_fin])
            else:
                fecha_inicio = proyecto.fecha_inicio

            if fecha_inicio:
                tarea.fecha_inicio = make_aware(fecha_inicio)
                tarea.fecha_fin = tarea.fecha_inicio + timedelta(days=tarea.duracion)
                tarea.save()

    return {"mensaje": "Tareas cargadas correctamente."}